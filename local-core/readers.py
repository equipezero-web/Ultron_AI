from __future__ import annotations

import csv
import json
import shutil
import subprocess
import wave
from pathlib import Path
from typing import Iterable

from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader
from pptx import Presentation

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".html", ".htm", ".css", ".scss", ".sass",
    ".js", ".mjs", ".cjs", ".ts", ".tsx", ".jsx", ".json", ".jsonc", ".xml",
    ".yml", ".yaml", ".toml", ".ini", ".cfg", ".conf", ".env.example", ".sql",
    ".py", ".java", ".kt", ".kts", ".c", ".h", ".cpp", ".hpp", ".cs", ".go",
    ".rs", ".php", ".rb", ".sh", ".bat", ".ps1", ".vue", ".svelte", ".rtf"
}

TABLE_EXTENSIONS = {".csv", ".tsv", ".xlsx", ".xlsm"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".pptx"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
AUDIO_EXTENSIONS = {".mp3", ".wav", ".ogg", ".m4a", ".flac", ".aac"}
VIDEO_EXTENSIONS = {".mp4", ".mov", ".mkv", ".webm", ".avi", ".m4v"}
MEDIA_EXTENSIONS = IMAGE_EXTENSIONS | AUDIO_EXTENSIONS | VIDEO_EXTENSIONS | {".svg"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | TABLE_EXTENSIONS | DOCUMENT_EXTENSIONS | MEDIA_EXTENSIONS


def _clean(text: str) -> str:
    return "\n".join(line.rstrip() for line in str(text or "").replace("\r\n", "\n").split("\n")).strip()


def _read_text(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return _clean(path.read_text(encoding=encoding, errors="strict"))
        except UnicodeDecodeError:
            continue
    return _clean(path.read_text(encoding="utf-8", errors="ignore"))


def _read_delimited(path: Path, delimiter: str) -> str:
    rows: list[str] = []
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                reader = csv.reader(handle, delimiter=delimiter)
                for row in reader:
                    if any(str(value).strip() for value in row):
                        rows.append("\t".join(str(value) for value in row))
            break
        except UnicodeDecodeError:
            rows = []
            continue
    return "[TABELA]\n" + "\n".join(rows) if rows else ""


def _read_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = _clean(page.extract_text() or "")
        if text:
            pages.append(f"[PÁGINA {index}]\n{text}")
    return "\n\n".join(pages)


def _read_docx(path: Path) -> str:
    document = Document(str(path))
    blocks = []
    for paragraph in document.paragraphs:
        text = _clean(paragraph.text)
        if text:
            blocks.append(text)
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                rows.append("\t".join(values))
        if rows:
            blocks.append(f"[TABELA {table_index}]\n" + "\n".join(rows))
    return "\n\n".join(blocks)


def _read_xlsx(path: Path) -> str:
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    blocks = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(value.strip() for value in values):
                    rows.append("\t".join(values))
            if rows:
                blocks.append(f"[PLANILHA: {sheet.title}]\n" + "\n".join(rows))
    finally:
        workbook.close()
    return "\n\n".join(blocks)


def _read_pptx(path: Path) -> str:
    presentation = Presentation(str(path))
    blocks = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = _clean(shape.text)
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False):
                rows = []
                for row in shape.table.rows:
                    values = [cell.text.strip() for cell in row.cells]
                    if any(values):
                        rows.append("\t".join(values))
                if rows:
                    parts.append("\n".join(rows))
        if parts:
            blocks.append(f"[SLIDE {slide_index}]\n" + "\n\n".join(parts))
    return "\n\n".join(blocks)


def _ffprobe_metadata(path: Path) -> dict:
    ffprobe = shutil.which("ffprobe")
    if not ffprobe:
        return {}
    try:
        completed = subprocess.run(
            [
                ffprobe,
                "-v", "quiet",
                "-print_format", "json",
                "-show_format",
                "-show_streams",
                str(path),
            ],
            check=False,
            capture_output=True,
            text=True,
            timeout=4,
        )
        if completed.returncode != 0 or not completed.stdout.strip():
            return {}
        return json.loads(completed.stdout)
    except Exception:
        return {}


def _image_metadata(path: Path, relative_path: str) -> str:
    stat = path.stat()
    lines = [
        "[ARQUIVO DE IMAGEM]",
        f"Nome: {path.name}",
        f"Caminho relativo: {relative_path}",
        f"Extensão: {path.suffix.lower()}",
        f"Tamanho: {stat.st_size} bytes",
    ]
    try:
        with Image.open(path) as image:
            lines.extend([
                f"Dimensões: {image.width}x{image.height}",
                f"Formato: {image.format or path.suffix.lstrip('.').upper()}",
                f"Modo de cor: {image.mode}",
            ])
            if getattr(image, "n_frames", 1) > 1:
                lines.append(f"Frames: {image.n_frames}")
    except Exception:
        pass
    lines.append("Conteúdo visual não é inferido localmente; o RAG indexa metadados seguros desta mídia.")
    return "\n".join(lines)


def _audio_metadata(path: Path, relative_path: str) -> str:
    stat = path.stat()
    lines = [
        "[ARQUIVO DE ÁUDIO]",
        f"Nome: {path.name}",
        f"Caminho relativo: {relative_path}",
        f"Extensão: {path.suffix.lower()}",
        f"Tamanho: {stat.st_size} bytes",
    ]
    if path.suffix.lower() == ".wav":
        try:
            with wave.open(str(path), "rb") as audio:
                duration = audio.getnframes() / max(1, audio.getframerate())
                lines.extend([
                    f"Duração: {duration:.2f}s",
                    f"Canais: {audio.getnchannels()}",
                    f"Sample rate: {audio.getframerate()} Hz",
                ])
        except Exception:
            pass
    probe = _ffprobe_metadata(path)
    fmt = probe.get("format") if isinstance(probe, dict) else None
    if isinstance(fmt, dict) and fmt.get("duration"):
        try:
            lines.append(f"Duração: {float(fmt['duration']):.2f}s")
        except Exception:
            pass
    lines.append("Áudio não é transcrito automaticamente nesta camada; metadados e nome permanecem pesquisáveis.")
    return "\n".join(dict.fromkeys(lines))


def _video_metadata(path: Path, relative_path: str) -> str:
    stat = path.stat()
    lines = [
        "[ARQUIVO DE VÍDEO]",
        f"Nome: {path.name}",
        f"Caminho relativo: {relative_path}",
        f"Extensão: {path.suffix.lower()}",
        f"Tamanho: {stat.st_size} bytes",
    ]
    probe = _ffprobe_metadata(path)
    if isinstance(probe, dict):
        fmt = probe.get("format") or {}
        streams = probe.get("streams") or []
        try:
            if fmt.get("duration"):
                lines.append(f"Duração: {float(fmt['duration']):.2f}s")
        except Exception:
            pass
        video = next((stream for stream in streams if stream.get("codec_type") == "video"), None)
        if video:
            if video.get("width") and video.get("height"):
                lines.append(f"Dimensões: {video['width']}x{video['height']}")
            if video.get("codec_name"):
                lines.append(f"Codec: {video['codec_name']}")
    lines.append("Vídeo não é interpretado quadro a quadro nesta camada; metadados e nome permanecem pesquisáveis.")
    return "\n".join(lines)


def _svg_metadata(path: Path, relative_path: str) -> str:
    text = _read_text(path)
    return f"[ARQUIVO SVG]\nCaminho relativo: {relative_path}\n{text[:120000]}"


def read_file(path: Path, root: Path) -> str:
    path = path.resolve()
    root = root.resolve()
    relative_path = path.relative_to(root).as_posix()
    suffix = path.suffix.lower()

    if suffix in TEXT_EXTENSIONS:
        return _read_text(path)
    if suffix == ".csv":
        return _read_delimited(path, ",")
    if suffix == ".tsv":
        return _read_delimited(path, "\t")
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx(path)
    if suffix == ".pptx":
        return _read_pptx(path)
    if suffix == ".svg":
        return _svg_metadata(path, relative_path)
    if suffix in IMAGE_EXTENSIONS:
        return _image_metadata(path, relative_path)
    if suffix in AUDIO_EXTENSIONS:
        return _audio_metadata(path, relative_path)
    if suffix in VIDEO_EXTENSIONS:
        return _video_metadata(path, relative_path)
    return ""


def chunk_text(text: str, *, chunk_size: int = 1400, overlap: int = 180) -> list[str]:
    text = _clean(text)
    if not text:
        return []

    chunk_size = max(400, int(chunk_size))
    overlap = max(0, min(int(overlap), chunk_size // 2))

    paragraphs = [item.strip() for item in text.split("\n\n") if item.strip()]
    chunks: list[str] = []
    buffer = ""

    for paragraph in paragraphs:
        candidate = f"{buffer}\n\n{paragraph}".strip() if buffer else paragraph
        if len(candidate) <= chunk_size:
            buffer = candidate
            continue

        if buffer:
            chunks.append(buffer)
            tail = buffer[-overlap:] if overlap else ""
            buffer = f"{tail}\n\n{paragraph}".strip()
        else:
            start = 0
            while start < len(paragraph):
                end = min(len(paragraph), start + chunk_size)
                chunks.append(paragraph[start:end])
                if end >= len(paragraph):
                    buffer = ""
                    break
                start = max(start + 1, end - overlap)

    if buffer:
        chunks.append(buffer)

    return [chunk.strip() for chunk in chunks if chunk.strip()]


def iter_supported_files(root: Path) -> Iterable[Path]:
    root = root.resolve()
    ignored_names = {".git", ".ultron_index", "__pycache__", ".venv", "node_modules"}
    for path in root.rglob("*"):
        if not path.is_file():
            continue
        if any(part in ignored_names for part in path.parts):
            continue
        if path.name.startswith("."):
            continue
        if path.suffix.lower() in SUPPORTED_EXTENSIONS:
            yield path
