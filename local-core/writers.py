from __future__ import annotations

import csv
import re
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

TEXT_WRITABLE = {
    ".txt", ".md", ".markdown", ".html", ".htm", ".css", ".scss", ".sass",
    ".js", ".mjs", ".cjs", ".ts", ".tsx", ".jsx", ".json", ".jsonc", ".xml",
    ".yml", ".yaml", ".toml", ".ini", ".cfg", ".conf", ".sql", ".py", ".java",
    ".kt", ".kts", ".c", ".h", ".cpp", ".hpp", ".cs", ".go", ".rs", ".php",
    ".rb", ".sh", ".bat", ".ps1", ".vue", ".svelte", ".rtf"
}

TABLE_WRITABLE = {".csv", ".tsv"}
BINARY_WRITABLE = {".docx", ".xlsx", ".pdf", ".pptx"}
WRITABLE_EXTENSIONS = TEXT_WRITABLE | TABLE_WRITABLE | BINARY_WRITABLE


def sanitize_filename(value: str) -> str:
    name = Path(str(value or "arquivo_ultron.txt")).name
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip(" .")
    return name or "arquivo_ultron.txt"


def safe_export_suffix(source_suffix: str, requested_suffix: str | None = None) -> str:
    requested = str(requested_suffix or "").lower()
    if requested in WRITABLE_EXTENSIONS:
        return requested
    source = str(source_suffix or "").lower()
    if source == ".xlsm":
        return ".xlsx"
    if source in WRITABLE_EXTENSIONS:
        return source
    return ".txt"


def _write_text(target: Path, content: str) -> None:
    target.write_text(str(content), encoding="utf-8")


def _write_delimited(target: Path, content: str, delimiter: str) -> None:
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle, delimiter=delimiter)
        for raw_line in str(content).splitlines():
            line = raw_line.strip("\n")
            if not line.strip() or line.strip() == "[TABELA]":
                continue
            if line.startswith("[PLANILHA:") and line.endswith("]"):
                continue
            if "\t" in line:
                values = line.split("\t")
            elif delimiter in line:
                values = next(csv.reader([line], delimiter=delimiter))
            else:
                values = [line]
            writer.writerow(values)


def _write_docx(target: Path, content: str) -> None:
    document = Document()
    blocks = [block.strip() for block in str(content).split("\n\n") if block.strip()]
    if not blocks:
        blocks = [""]
    for block in blocks:
        lines = block.splitlines()
        if len(lines) > 1 and lines[0].startswith("[") and lines[0].endswith("]"):
            document.add_heading(lines[0].strip("[]"), level=2)
            body = "\n".join(lines[1:]).strip()
            if body:
                document.add_paragraph(body)
        else:
            document.add_paragraph(block)
    document.save(str(target))


def _write_xlsx(target: Path, content: str) -> None:
    workbook = Workbook()
    default = workbook.active
    default.title = "Ultron"
    current = default
    created = 1

    for raw_line in str(content).splitlines():
        line = raw_line.rstrip("\n")
        if line.startswith("[PLANILHA:") and line.endswith("]"):
            title = line[len("[PLANILHA:"):-1].strip() or f"Planilha{created + 1}"
            if current.max_row == 1 and current.cell(1, 1).value is None and created == 1:
                current.title = title[:31]
            else:
                current = workbook.create_sheet(title=title[:31])
                created += 1
            continue
        if not line.strip() or line.strip() == "[TABELA]":
            continue
        values = line.split("\t") if "\t" in line else [line]
        current.append(values)

    workbook.save(str(target))


def _write_pdf(target: Path, content: str) -> None:
    styles = getSampleStyleSheet()
    story = []
    for block in [item.strip() for item in str(content).split("\n\n") if item.strip()]:
        escaped = (
            block.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace("\n", "<br/>")
        )
        story.append(Paragraph(escaped, styles["BodyText"]))
        story.append(Spacer(1, 8))
    document = SimpleDocTemplate(str(target), pagesize=A4, title=target.stem)
    document.build(story or [Paragraph("", styles["BodyText"])])


def _write_pptx(target: Path, content: str) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    blocks = [item.strip() for item in str(content).split("\n\n") if item.strip()]
    if not blocks:
        blocks = ["Ultron"]

    for index, block in enumerate(blocks):
        layout = presentation.slide_layouts[1]
        slide = presentation.slides.add_slide(layout)
        lines = block.splitlines()
        title = lines[0]
        if title.startswith("[SLIDE") and title.endswith("]"):
            title = title.strip("[]")
            body = "\n".join(lines[1:]).strip()
        else:
            title = f"Ultron {index + 1}"
            body = block
        slide.shapes.title.text = title[:120]
        placeholder = slide.placeholders[1]
        placeholder.text = body[:7000]

    presentation.save(str(target))


def export_content(target: Path, content: str) -> str:
    suffix = target.suffix.lower()
    target.parent.mkdir(parents=True, exist_ok=True)

    if suffix in TEXT_WRITABLE:
        _write_text(target, content)
    elif suffix == ".csv":
        _write_delimited(target, content, ",")
    elif suffix == ".tsv":
        _write_delimited(target, content, "\t")
    elif suffix == ".docx":
        _write_docx(target, content)
    elif suffix == ".xlsx":
        _write_xlsx(target, content)
    elif suffix == ".pdf":
        _write_pdf(target, content)
    elif suffix == ".pptx":
        _write_pptx(target, content)
    else:
        raise ValueError(f"Formato de exportação não suportado: {suffix}")

    return suffix
