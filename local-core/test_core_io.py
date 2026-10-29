from __future__ import annotations

import csv
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from readers import read_file
from writers import export_content


class LocalCoreIoTests(unittest.TestCase):
    def test_text_and_csv_readers(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            text = root / "notes.md"
            text.write_text("Ultron knowledge base", encoding="utf-8")
            csv_file = root / "table.csv"
            csv_file.write_text("name,value\nalpha,10\n", encoding="utf-8")

            self.assertIn("Ultron knowledge base", read_file(text, root))
            csv_text = read_file(csv_file, root)
            self.assertIn("name\tvalue", csv_text)
            self.assertIn("alpha\t10", csv_text)

    def test_csv_and_tsv_writers_create_real_tables(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            csv_target = root / "result.csv"
            tsv_target = root / "result.tsv"
            content = "[TABELA]\nNome\tValor\nTeste\t42"

            export_content(csv_target, content)
            export_content(tsv_target, content)

            with csv_target.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.reader(handle))
            self.assertEqual(rows, [["Nome", "Valor"], ["Teste", "42"]])

            with tsv_target.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.reader(handle, delimiter="\t"))
            self.assertEqual(rows, [["Nome", "Valor"], ["Teste", "42"]])

    def test_docx_reader_and_writer(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            target = root / "result.docx"
            export_content(target, "Título\n\nConteúdo validado pelo Ultron")
            document = Document(str(target))
            combined = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("Conteúdo validado pelo Ultron", combined)
            self.assertIn("Conteúdo validado pelo Ultron", read_file(target, root))

    def test_xlsx_reader_and_writer(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            target = root / "result.xlsx"
            export_content(target, "[PLANILHA: Dados]\nNome\tValor\nTeste\t42")
            workbook = load_workbook(str(target), read_only=True, data_only=True)
            try:
                sheet = workbook.worksheets[0]
                self.assertEqual(sheet.title, "Dados")
                self.assertEqual(sheet["A2"].value, "Nome")
                self.assertEqual(sheet["B2"].value, "Valor")
                self.assertEqual(sheet["A3"].value, "Teste")
                self.assertEqual(sheet["B3"].value, "42")
            finally:
                workbook.close()
            self.assertIn("Teste\t42", read_file(target, root))

    def test_pptx_reader_and_writer(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            target = root / "result.pptx"
            export_content(target, "[SLIDE 1]\nArquitetura Local RAG\n\n[SLIDE 2]\nBusca híbrida")
            presentation = Presentation(str(target))
            self.assertGreaterEqual(len(presentation.slides), 2)
            self.assertIn("Arquitetura Local RAG", read_file(target, root))

    def test_source_files_are_not_overwritten_by_exporters(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "source.txt"
            target = root / "export.txt"
            source.write_text("ORIGINAL", encoding="utf-8")
            export_content(target, "NOVA VERSÃO")
            self.assertEqual(source.read_text(encoding="utf-8"), "ORIGINAL")
            self.assertEqual(target.read_text(encoding="utf-8"), "NOVA VERSÃO")


if __name__ == "__main__":
    unittest.main()
